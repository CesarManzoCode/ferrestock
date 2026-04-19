import io
from typing import Any

import openpyxl
from rapidfuzz import fuzz

# Aliases para autodetección de columnas
CAMPO_ALIASES: dict[str, list[str]] = {
    "nombre": ["nombre", "producto", "descripcion", "articulo", "item", "name"],
    "precio": ["precio", "price", "costo", "p.u.", "precio unitario", "valor"],
    "existencias": ["existencias", "stock", "cantidad", "inventario", "piezas", "qty"],
    "descripcion": ["descripcion larga", "detalle", "notas", "description"],
}

SCORE_AUTODETECCION = 75  # umbral para considerar una columna detectada automáticamente


def detectar_columna(nombre_columna: str) -> str | None:
    """Dado el nombre de una columna del Excel, retorna el campo de FerreStock
    que mejor coincide, o None si no hay coincidencia suficiente."""
    nombre_lower = nombre_columna.lower().strip()

    mejor_campo = None
    mejor_score = 0

    for campo, aliases in CAMPO_ALIASES.items():
        for alias in aliases:
            score = fuzz.ratio(nombre_lower, alias)
            if score > mejor_score:
                mejor_score = score
                mejor_campo = campo

    return mejor_campo if mejor_score >= SCORE_AUTODETECCION else None


def leer_columnas_excel(
    contenido: bytes,
    fila_inicio: int = 1,
) -> dict:
    """
    Lee el Excel y retorna:
    - columnas: lista de nombres de columna encontrados en fila_inicio
    - mapeo_sugerido: dict columna -> campo detectado automáticamente
    - preview: primeras 5 filas de datos para que el usuario confirme
    """
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    ws = wb.active

    filas = list(ws.iter_rows(min_row=fila_inicio, values_only=True))

    if not filas:
        return {"columnas": [], "mapeo_sugerido": {}, "preview": []}

    encabezados = [
        str(celda).strip() if celda is not None else f"Columna_{i+1}"
        for i, celda in enumerate(filas[0])
    ]

    mapeo_sugerido = {}
    for col in encabezados:
        campo = detectar_columna(col)
        if campo:
            mapeo_sugerido[col] = campo

    # Preview de hasta 5 filas de datos (sin el encabezado)
    preview = []
    for fila in filas[1:6]:
        preview.append({
            encabezados[i]: celda
            for i, celda in enumerate(fila)
            if i < len(encabezados)
        })

    wb.close()

    return {
        "columnas": encabezados,
        "mapeo_sugerido": mapeo_sugerido,
        "preview": preview,
    }


def importar_productos_desde_excel(
    contenido: bytes,
    mapeo: dict[str, str],        # {nombre_columna_excel: campo_ferrestock}
    campos_extra_cols: list[str], # columnas que se guardan como campo extra
    fila_inicio: int = 1,
    tenant_id: int = 0,
) -> tuple[list[dict[str, Any]], list[dict]]:
    """
    Procesa el Excel con el mapeo confirmado por el usuario.
    Retorna:
    - productos: lista de dicts listos para insertar en DB
    - errores: filas que no pudieron procesarse con su motivo
    """
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    ws = wb.active

    filas = list(ws.iter_rows(min_row=fila_inicio, values_only=True))
    if not filas:
        wb.close()
        return [], []

    encabezados = [
        str(c).strip() if c is not None else f"Columna_{i+1}"
        for i, c in enumerate(filas[0])
    ]

    productos = []
    errores = []

    for num_fila, fila in enumerate(filas[1:], start=fila_inicio + 1):
        fila_dict = {
            encabezados[i]: celda
            for i, celda in enumerate(fila)
            if i < len(encabezados)
        }

        try:
            nombre_col = next(
                (col for col, campo in mapeo.items() if campo == "nombre"), None
            )
            precio_col = next(
                (col for col, campo in mapeo.items() if campo == "precio"), None
            )

            nombre = str(fila_dict.get(nombre_col, "")).strip() if nombre_col else ""
            precio_raw = fila_dict.get(precio_col, 0) if precio_col else 0

            if not nombre:
                errores.append({"fila": num_fila, "motivo": "Nombre vacío", "datos": fila_dict})
                continue

            try:
                precio = float(str(precio_raw).replace(",", "").replace("$", "").strip())
            except (ValueError, TypeError):
                errores.append({"fila": num_fila, "motivo": f"Precio inválido: {precio_raw}", "datos": fila_dict})
                continue

            # Campos opcionales fijos
            existencias = 0
            existencias_col = next(
                (col for col, campo in mapeo.items() if campo == "existencias"), None
            )
            if existencias_col and fila_dict.get(existencias_col) is not None:
                try:
                    existencias = float(str(fila_dict[existencias_col]).replace(",", "").strip())
                except (ValueError, TypeError):
                    existencias = 0

            descripcion = None
            descripcion_col = next(
                (col for col, campo in mapeo.items() if campo == "descripcion"), None
            )
            if descripcion_col:
                descripcion = str(fila_dict.get(descripcion_col, "") or "").strip() or None

            # Campos extra — usar el nombre_campo del mapeo, no el nombre de columna
            campos_extra = {}
            for col in campos_extra_cols:
                val = fila_dict.get(col)
                if val is not None:
                    # mapeo[col] = nombre_campo (ej: "precio_iva"), si no existe usar col
                    clave = mapeo.get(col, col)
                    campos_extra[clave] = val

            productos.append({
                "tenant_id": tenant_id,
                "nombre": nombre,
                "precio": precio,
                "existencias": existencias,
                "descripcion": descripcion,
                "campos_extra": campos_extra if campos_extra else None,
            })

        except Exception as e:
            errores.append({"fila": num_fila, "motivo": str(e), "datos": fila_dict})

    wb.close()
    return productos, errores
